"""Gera docs/folha_de_teste.docx a partir de docs/folha_de_teste.md.

Uso: python tools/gerar_folha_docx.py   (precisa de: pip install python-docx)
O .md é a fonte. Edite o .md e rode de novo para atualizar o .docx.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

RAIZ = Path(__file__).resolve().parent.parent
ORIGEM = RAIZ / "docs" / "folha_de_teste.md"
DESTINO = RAIZ / "docs" / "folha_de_teste.docx"


def limpa(txt):
    """Tira a marcação markdown dos trechos curtos."""
    return txt.replace("`", "").replace("**", "")


def runs(par, txt):
    """Escreve o texto, com negrito para **trecho**."""
    for i, parte in enumerate(re.split(r"\*\*(.+?)\*\*", txt)):
        if parte:
            r = par.add_run(parte.replace("`", ""))
            r.bold = i % 2 == 1


def sombra(celula, cor):
    tc = celula._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), cor)
    tc.append(sh)


def tabela(doc, linhas):
    cab = [c.strip() for c in linhas[0].strip("|").split("|")]
    corpo = [[c.strip() for c in l.strip("|").split("|")] for l in linhas[2:]]
    t = doc.add_table(rows=1, cols=len(cab))
    t.style = "Table Grid"
    for i, c in enumerate(cab):
        cel = t.rows[0].cells[i]
        cel.text = ""
        r = cel.paragraphs[0].add_run(limpa(c))
        r.bold = True
        sombra(cel, "D9E2F3")
    # colunas de anotação precisam de espaço para escrever à mão
    livre = {i for i, c in enumerate(cab) if c.lower() in ("o que não bateu", "nome real", "tempo real", "ok / aj / x", "existe?", "resposta")}
    for lin in corpo:
        cels = t.add_row().cells
        for i, c in enumerate(lin[: len(cab)]):
            cels[i].text = ""
            runs(cels[i].paragraphs[0], c)
        if livre:
            t.rows[-1].height = Cm(1.3)
    for row in t.rows:
        for cel in row.cells:
            for p in cel.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Cm(1.6))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    linhas = ORIGEM.read_text(encoding="utf8").splitlines()
    i = 0
    while i < len(linhas):
        l = linhas[i]
        if l.startswith("# "):
            doc.add_heading(limpa(l[2:]), 0)
        elif l.startswith("## "):
            doc.add_heading(limpa(l[3:]), 1)
        elif l.startswith("|"):
            bloco = []
            while i < len(linhas) and linhas[i].startswith("|"):
                bloco.append(linhas[i])
                i += 1
            tabela(doc, bloco)
            continue
        elif l.startswith("- [ ] "):
            runs(doc.add_paragraph(), "☐  " + l[6:])
        elif re.match(r"^\d+\. ", l):
            runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", l))
        elif l.startswith("- "):
            runs(doc.add_paragraph(style="List Bullet"), l[2:])
        elif l.strip():
            runs(doc.add_paragraph(), l)
        i += 1
    doc.save(DESTINO)
    print("Gerado:", DESTINO)


if __name__ == "__main__":
    main()
